'''Generate Viva 1 PowerPoint from hardcoded slide content.'''

from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt


SLIDES = [
    {
        'title': 'SENTINEL-AI',
        'subtitle': 'Pose-Estimation & Spatial Graph Analytics for Real-Time Violence Detection in CCTV Feeds',
        'bullets': [
            'Team: [names]',
            'Mentor: [name]',
            'Domain: Computer Vision · Distributed Systems · Public Safety Tech'
        ]
    },
    {
        'title': 'Problem Statement',
        'bullets': [
            'Cities have millions of CCTV cameras but human operators miss ~95% of incidents.',
            'Average response time to fights/assault is 8–15 minutes.',
            'Manual monitoring is not scalable and suffers from fatigue-induced errors.'
        ]
    },
    {
        'title': 'Proposed Solution',
        'bullets': [
            'Replace heavy 3D-CNNs with lightweight 2D pose estimation (17 keypoints per person).',
            'Explainable heuristics: joint velocity + proximity + bounding-box overlap.',
            'End-to-end pipeline: ingest → pose → analytics → alerts → dashboard.'
        ]
    },
    {
        'title': 'Literature Survey & Research Gaps',
        'bullets': [
            'Computational Overhead Gap: 3D-CNNs/ViTs need GPUs and run >150 ms/frame on CPU.',
            'False-Positive Gap: Optical flow confuses running, dancing, or hugging with violence.',
            'Integration Gap: Most work stops at offline benchmarks; streaming/concurrency/alerting are rarely addressed.'
        ]
    },
    {
        'title': 'System Architecture',
        'bullets': [
            'Ingest Layer (OpenCV) → Pose Estimator (YOLOv8-Pose) → Analytics Engine',
            'Analytics: joint velocity, k-d Tree proximity, IoU overlap, threat scoring',
            'SQLite/PostgreSQL incident DB, Redis/Celery async alerts, FastAPI/WebSockets, Next.js dashboard'
        ]
    },
    {
        'title': 'Methodology',
        'bullets': [
            'Pose Extraction: YOLOv8-Pose detects persons and 17 keypoints.',
            'Motion Analysis: frame-to-frame velocity and acceleration of wrists, elbows, head.',
            'Proximity Analysis: distance between person centroids per frame.',
            'Threat Scoring: Phase 2 adds k-d Tree + IoU + temporal window.'
        ]
    },
    {
        'title': 'Viva 1 Deliverables',
        'bullets': [
            'Repository scaffolded with modules and docs.',
            'Working CLI demo on normal and fight test videos.',
            'Annotated output video with skeletal overlay and bounding boxes.',
            'pytest unit tests for motion and proximity math.',
            'SRS, literature survey, research gaps, and architecture documents.'
        ]
    },
    {
        'title': 'Demo Script',
        'bullets': [
            'Run CLI demo on a normal clip; show NORMAL flags.',
            'Run CLI demo on a fight clip; show high velocity values and FIGHT flags.',
            'Show live webcam overlay briefly to prove real-time feasibility.'
        ]
    },
    {
        'title': 'Likely Questions & Answers',
        'bullets': [
            'Why not 3D-CNN? Needs GPU, >150 ms/frame. Our pose approach reduces dimensionality by 90% and runs on CPU.',
            'Hug vs fight? Hugging has low kinetic variance; fights have rapid multi-directional limb acceleration.',
            'Role of k-d Trees? O(N log N) neighbor queries vs naive O(N2) proximity checks.',
            'How avoid freezing video? Redis + Celery async alert dispatch in Phase 3.'
        ]
    },
    {
        'title': 'Roadmap',
        'bullets': [
            'Phase 1 (Viva 1): Pose extraction + CLI demo + docs',
            'Phase 2 (Viva 2): k-d Tree, IoU, temporal window, benchmarks',
            'Phase 3 (Viva 3): Threading, FastAPI/WebSockets, React UI, alerts',
            'Phase 4 (Viva 4): Dataset evaluation, Docker, report, paper draft'
        ]
    }
]


def add_title_slide(prs, slide):
    layout = prs.slide_layouts[0]
    s = prs.slides.add_slide(layout)
    s.shapes.title.text = slide['title']
    tf = s.placeholders[1].text_frame
    tf.clear()
    tf.text = slide['subtitle']
    for b in slide['bullets']:
        p = tf.add_paragraph()
        p.text = b
        p.level = 0
        p.font.size = Pt(18)


def add_bullet_slide(prs, slide):
    layout = prs.slide_layouts[1]
    s = prs.slides.add_slide(layout)
    s.shapes.title.text = slide['title']
    tf = s.shapes.placeholders[1].text_frame
    tf.text = ''
    for b in slide['bullets']:
        p = tf.add_paragraph()
        p.text = b
        p.level = 0
        p.font.size = Pt(20)


def main():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    for i, slide in enumerate(SLIDES):
        if i == 0:
            add_title_slide(prs, slide)
        else:
            add_bullet_slide(prs, slide)

    output = Path('docs/viva1/SENTINEL_AI_Viva1.pptx')
    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output))
    print('Saved {}'.format(output))


if __name__ == '__main__':
    main()
